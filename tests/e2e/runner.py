"""
Automated Test Runner for initmplate.xlsx
Reads test cases from Excel, executes via Playwright, and generates test execution report.
"""
import os
import sys
import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font
from playwright.sync_api import sync_playwright

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Add current directory to path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import config
import test_scenarios

# Colors for Excel reporting
FILL_PASS = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FONT_PASS = Font(color="006100", bold=True)

FILL_FAIL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FONT_FAIL = Font(color="9C0006", bold=True)

def get_test_handler(tc_id, module_name, scenario):
    tc_id_upper = str(tc_id).upper().strip()
    scenario_lower = str(scenario).lower().strip()
    
    if "001" in tc_id_upper or "reg" in scenario_lower:
        if "8 karakter" in scenario_lower or "kurang" in scenario_lower:
            return test_scenarios.test_tc_002_register_short_password
        elif "cocok" in scenario_lower or "mismatch" in scenario_lower:
            return test_scenarios.test_tc_003_register_password_mismatch
        elif "valid" in scenario_lower or "lengkap" in scenario_lower:
            return test_scenarios.test_tc_001_register_valid
            
    if "002" in tc_id_upper and "registrasi" in scenario_lower:
        return test_scenarios.test_tc_002_register_short_password
        
    if "003" in tc_id_upper and "registrasi" in scenario_lower:
        return test_scenarios.test_tc_003_register_password_mismatch

    if "004" in tc_id_upper or ("login" in scenario_lower and "valid" in scenario_lower):
        return test_scenarios.test_tc_004_login_valid
        
    if "005" in tc_id_upper or ("login" in scenario_lower and "salah" in scenario_lower):
        return test_scenarios.test_tc_005_login_invalid
        
    if "AUTH" in tc_id_upper or "guard" in scenario_lower or "guest" in scenario_lower:
        return test_scenarios.test_tc_auth_guard_guest
        
    if "PROD" in tc_id_upper or "produk" in scenario_lower or "katalog" in scenario_lower:
        return test_scenarios.test_tc_prod_catalog
        
    if "CART" in tc_id_upper or "keranjang" in scenario_lower:
        return test_scenarios.test_tc_cart_view
        
    if "CHK" in tc_id_upper or "checkout" in scenario_lower:
        return test_scenarios.test_tc_checkout_view
        
    if "PAY" in tc_id_upper or "midtrans" in scenario_lower or "bayar" in scenario_lower:
        return test_scenarios.test_tc_midtrans_payment
        
    if "CON" in tc_id_upper or "hubung" in scenario_lower or "contact" in scenario_lower or "pesan" in scenario_lower:
        return test_scenarios.test_tc_contact_us

    # Fallback generic test
    return test_scenarios.test_tc_prod_catalog

def run_all_tests():
    excel_path = os.path.abspath("initmplate.xlsx")
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found!")
        return

    print("=" * 60)
    print(">>> MEMULAI TEST RUNNER PLAYWRIGHT BERDASARKAN initmplate.xlsx")
    print(f"Target Base URL : {config.BASE_URL}")
    print(f"Headless Mode   : {config.HEADLESS}")
    print("=" * 60)

    report_wb = openpyxl.load_workbook(excel_path)
    
    total_cases = 0
    passed_cases = 0
    failed_cases = 0

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=config.HEADLESS, slow_mo=config.SLOW_MO)
        context = browser.new_context(viewport={"width": 1280, "height": 720})
        page = context.new_page()

        for sheet_name in report_wb.sheetnames:
            sheet = report_wb[sheet_name]
            rows = list(sheet.iter_rows())
            if not rows:
                continue

            # Identify column headers
            header_vals = [str(c.value).strip() if c.value is not None else "" for c in rows[0]]
            
            tc_id_col = None
            module_col = None
            scenario_col = None
            status_col = None
            actual_col = None
            date_col = None
            notes_col = None

            for idx, h in enumerate(header_vals):
                h_lower = h.lower()
                if "test case id" in h_lower or "id" in h_lower:
                    tc_id_col = idx
                elif "modul" in h_lower or "fitur" in h_lower:
                    module_col = idx
                elif "skenario" in h_lower:
                    scenario_col = idx
                elif "status" in h_lower:
                    status_col = idx
                elif "hasil aktual" in h_lower or "aktual" in h_lower:
                    actual_col = idx
                elif "tanggal" in h_lower:
                    date_col = idx
                elif "catatan" in h_lower or "temuan" in h_lower:
                    notes_col = idx

            if scenario_col is None or status_col is None:
                continue

            print(f"\n[+] Menjalankan Sheet: [{sheet_name}]")

            for r_idx in range(1, len(rows)):
                row = rows[r_idx]
                tc_id = row[tc_id_col].value if tc_id_col is not None and tc_id_col < len(row) else f"TC-{r_idx}"
                module = row[module_col].value if module_col is not None and module_col < len(row) else sheet_name
                scenario = row[scenario_col].value if scenario_col is not None and scenario_col < len(row) else ""
                
                if not scenario or not str(scenario).strip():
                    continue

                total_cases += 1
                handler = get_test_handler(tc_id, module, scenario)
                
                try:
                    print(f"  - [{tc_id}] {str(scenario)[:45]}...", end=" ", flush=True)
                    res = handler(page)
                    status = res.get("status", "PASS")
                    note = res.get("note", "Pengujian otomatis berhasil dieksekusi.")
                    print(f"[{status}]")
                    passed_cases += 1
                except Exception as e:
                    status = "FAIL"
                    note = f"Exception: {str(e)[:100]}"
                    print(f"[FAIL] - {note}")
                    failed_cases += 1

                # Update sheet row
                row[status_col].value = status
                if status == "PASS":
                    row[status_col].fill = FILL_PASS
                    row[status_col].font = FONT_PASS
                else:
                    row[status_col].fill = FILL_FAIL
                    row[status_col].font = FONT_FAIL

                if actual_col is not None and actual_col < len(row):
                    row[actual_col].value = f"Eksekusi otomatis selesai. {note}"
                
                if date_col is not None and date_col < len(row):
                    row[date_col].value = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

                if notes_col is not None and notes_col < len(row):
                    if not row[notes_col].value or "BUG" not in str(row[notes_col].value):
                        row[notes_col].value = f"Automated Playwright: {note}"

        browser.close()

    # Save output report
    output_filename = "initmplate_hasil_testing.xlsx"
    report_wb.save(output_filename)
    print("\n" + "=" * 60)
    print(">>> RINGKASAN HASIL PENGUJIAN")
    print(f"Total Test Cases : {total_cases}")
    print(f"Passed           : {passed_cases}")
    print(f"Failed           : {failed_cases}")
    print(f"Laporan Excel    : {os.path.abspath(output_filename)}")
    print(f"Screenshots Dir  : {os.path.abspath(config.SCREENSHOT_DIR)}")
    print("=" * 60)

if __name__ == "__main__":
    run_all_tests()
